from io import BytesIO
from typing import Dict, List, Optional, Set, Tuple

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import and_, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.plugins.vehicle_plugin import models, schemas

# ============================================================
# Excel 中文表头 → Pydantic 英文字段名映射
# ============================================================
EXCEL_FIELD_MAPPING: Dict[str, str] = {
    "车型": "model",
    "组别": "group",
    "车辆阶段": "vehicle_stage",
    "车辆配置": "configuration",
    "车主权限": "owner_name",
    "车辆编号": "vehicle_code",
    "停车地点": "parking_location",
    "使用状态": "vehicle_status",
    "车辆状态": "test_status",
    "备注": "remarks",
    "VIN码": "vin_code",
    "驱动电机号/发动机号": "engine_num",
    "车牌号": "plate_number",
    "临牌到期时间": "temp_plate_expire_date",
    "临牌已办理次数": "temp_plate_count",
}

# 重复判定字段（vin_code + vehicle_code 组合唯一）
REPEAT_VEHICLE_FIELDS = ("vin_code", "vehicle_code")

# 需要强制转为字符串的字段（Excel 中纯数字列会被 openpyxl 读取为 int/float）
STRING_FIELDS = (
    "model",
    "vehicle_code",
    "vin_code",
    "owner_name",
    "plate_number",
    "vehicle_stage",
    "configuration",
    "parking_location",
    "engine_num",
    "remarks"
)

# ============================================================
# 预构建枚举值查找表（模块加载时一次性完成，避免每次请求重复遍历 __members__）
# 将 Excel 中可能出现的值（英文名 / 中文值）统一映射到数据库存储的中文值
# ============================================================
def _build_enum_lookup(enum_cls) -> Dict[str, str]:
    lookup: Dict[str, str] = {}
    for member in enum_cls:
        lookup[member.name.upper()] = member.value
        lookup[member.value.upper()] = member.value
    return lookup


VEHICLE_STATUS_LOOKUP = _build_enum_lookup(models.VehicleStatus)
VEHICLE_GROUP_LOOKUP = _build_enum_lookup(models.VehicleGroup)
TEST_STATUS_LOOKUP = _build_enum_lookup(models.TestStatus)

ENUM_FIELD_LOOKUPS = {
    "vehicle_status": VEHICLE_STATUS_LOOKUP,
    "group": VEHICLE_GROUP_LOOKUP,
    "test_status": TEST_STATUS_LOOKUP,
}


# ============================================================
# 工具函数
# ============================================================
def _is_empty_row(row_cells) -> bool:
    for cell in row_cells:
        v = cell.value
        if v is None:
            continue
        if isinstance(v, str) and v.strip() == "":
            continue
        return False
    return True


def _read_excel_from_upload(file: UploadFile) -> Tuple[List[dict], List[str]]:
    file_content = file.file.read()
    if not file_content:
        raise ValueError("上传的文件内容为空")

    wb = load_workbook(BytesIO(file_content), read_only=True)

    if not wb.sheetnames:
        wb.close()
        raise ValueError("Excel 文件中没有工作表")

    ws = wb[wb.sheetnames[0]]

    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = []
    header_count: Dict[str, int] = {}
    for idx, h in enumerate(raw_headers):
        if h is None:
            h = f"unknown_column_{idx + 1}"
        if h in header_count:
            header_count[h] += 1
            h = f"{h}_{header_count[h]}"
        else:
            header_count[h] = 0
        headers.append(h)

    dict_list: List[dict] = []
    for row in ws.iter_rows(min_row=2):
        if _is_empty_row(row):
            continue
        row_data = {}
        for idx, header in enumerate(headers):
            row_data[header] = row[idx].value if idx < len(row) else None
        dict_list.append(row_data)

    wb.close()
    return dict_list, headers


def _make_unique_key(data_dict: dict) -> Tuple:
    return tuple(
        (data_dict.get(f) or "").strip() if isinstance(data_dict.get(f), str) else data_dict.get(f)
        for f in REPEAT_VEHICLE_FIELDS
    )


def _transform_record(record: dict) -> Optional[dict]:
    transformed = {}

    for cn_key, en_key in EXCEL_FIELD_MAPPING.items():
        if cn_key in record:
            transformed[en_key] = record[cn_key]

    for field in STRING_FIELDS:
        if field in transformed and transformed[field] is not None:
            if not isinstance(transformed[field], str):
                transformed[field] = str(transformed[field])

    for enum_field, lookup in ENUM_FIELD_LOOKUPS.items():
        if enum_field in transformed and transformed[enum_field] is not None:
            raw = str(transformed[enum_field]).strip().upper()
            if raw in lookup:
                transformed[enum_field] = lookup[raw]

    return transformed


# ============================================================
# 核心入口：处理上传的车辆 Excel 文件并存入数据库
# ============================================================
async def process_vehicle_excel_upload(file: UploadFile, db: AsyncSession) -> dict:
    # 1. 校验上传
    if file is None:
        return {"code": 400, "message": "文件上传失败：未接收到文件", "data": None}
    if not file.filename:
        return {"code": 400, "message": "文件上传失败：文件名为空", "data": None}

    filename_lower = file.filename.lower()
    if not (filename_lower.endswith(".xlsx") or filename_lower.endswith(".xls")):
        return {
            "code": 400,
            "message": "文件上传失败：仅支持 .xlsx 或 .xls 格式",
            "data": None,
        }

    # 2. 读取 Excel
    try:
        raw_records, _ = _read_excel_from_upload(file)
    except ValueError as e:
        return {"code": 400, "message": f"文件上传失败：{e}", "data": None}
    except Exception as e:
        return {"code": 400, "message": f"文件上传失败：文件读取异常 - {e}", "data": None}

    total_excel_rows = len(raw_records)
    if total_excel_rows == 0:
        return {
            "code": 200,
            "message": "文件中没有有效数据行",
            "data": {
                "msg": "无有效数据可导入",
                "total_excel_rows": 0,
                "excel_duplicate_rows": 0,
                "db_duplicate_rows": 0,
                "success_import_rows": 0,
                "failed_import_rows": 0,
            },
        }

    # 3. 单次遍历完成：中文表头映射 + 类型转换 + 枚举转换 + 内存去重
    transformed_records: List[dict] = []
    seen_keys: Set[Tuple] = set()
    excel_duplicate_count = 0

    for record in raw_records:
        transformed = _transform_record(record)
        if transformed is None:
            continue
        unique_key = _make_unique_key(transformed)
        if unique_key in seen_keys:
            excel_duplicate_count += 1
            continue
        seen_keys.add(unique_key)
        transformed_records.append(transformed)

    if not transformed_records:
        return {
            "code": 200,
            "message": "导入完成（所有数据均为 Excel 内部重复）",
            "data": {
                "msg": "无有效数据可导入",
                "total_excel_rows": total_excel_rows,
                "excel_duplicate_rows": excel_duplicate_count,
                "db_duplicate_rows": 0,
                "success_import_rows": 0,
                "failed_import_rows": 0,
            },
        }

    # 4. 数据库去重：一次批量查询，避免 N+1
    query_conditions = []
    for record in transformed_records:
        field_conds = []
        for field in REPEAT_VEHICLE_FIELDS:
            value = record.get(field)
            if isinstance(value, str):
                value = value.strip()
            if value is None:
                field_conds.append(getattr(models.Vehicle, field).is_(None))
            else:
                field_conds.append(getattr(models.Vehicle, field) == value)
        query_conditions.append(and_(*field_conds))

    stmt = select(models.Vehicle).filter(or_(*query_conditions))
    result = await db.execute(stmt)
    existing_records = result.scalars().all()

    existing_keys: Set[Tuple] = {
        _make_unique_key({f: getattr(r, f) for f in REPEAT_VEHICLE_FIELDS})
        for r in existing_records
    }

    final_records: List[dict] = []
    db_duplicate_count = 0
    for record in transformed_records:
        if _make_unique_key(record) in existing_keys:
            db_duplicate_count += 1
            continue
        final_records.append(record)

    if not final_records:
        return {
            "code": 200,
            "message": "导入完成（所有数据已在数据库中）",
            "data": {
                "msg": "所有数据已存在，无需导入",
                "total_excel_rows": total_excel_rows,
                "excel_duplicate_rows": excel_duplicate_count,
                "db_duplicate_rows": db_duplicate_count,
                "success_import_rows": 0,
                "failed_import_rows": 0,
            },
        }

    # 5. Pydantic 校验
    valid_records: List[models.Vehicle] = []
    for idx, item_dict in enumerate(final_records):
        try:
            validated = schemas.VehicleCreate(**item_dict)
            valid_records.append(models.Vehicle(**validated.model_dump()))
        except Exception as e:
            return {
                "code": 400,
                "message": f"第 {idx + 2} 行数据格式错误：{e}",
                "data": None,
            }

    # 6. 批量写入：单次 add_all + commit，事务安全
    try:
        db.add_all(valid_records)
        await db.commit()
    except Exception as e:
        await db.rollback()
        return {"code": 400, "message": f"数据库写入失败：{e}", "data": None}

    # 7. 返回结果
    return {
        "code": 200,
        "message": "文件上传并导入成功",
        "data": {
            "msg": "批量导入完成",
            "total_excel_rows": total_excel_rows,
            "excel_duplicate_rows": excel_duplicate_count,
            "db_duplicate_rows": db_duplicate_count,
            "success_import_rows": len(valid_records),
            "failed_import_rows": 0,
        },
    }